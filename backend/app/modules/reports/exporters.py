"""
Exportadores reais de tabela: CSV (UTF-8 com BOM, separador ``;`` — padrão do Excel
pt-BR), XLSX (openpyxl) e PDF (reportlab, via ``app.modules.reports.pdf``).

Proteção contra injeção de fórmula (CSV/XLSX): células de texto iniciadas por
``= + - @`` (ou TAB/CR) recebem o prefixo ``'``.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Optional, Sequence

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


@dataclass
class Column:
    key: str
    label: str
    kind: str = "text"  # text | int | money | date | datetime | bps


@dataclass
class Dataset:
    title: str
    columns: list[Column]
    rows: list[dict]
    meta: list[tuple[str, str]] = field(default_factory=list)  # filtros/período/observações
    notes: list[str] = field(default_factory=list)


def _safe_text(v: str) -> str:
    return "'" + v if v.startswith(FORMULA_PREFIXES) else v


def _display(col: Column, v: Any) -> Any:
    from app.core.money import format_brl

    if v is None:
        return ""
    if col.kind == "money" and isinstance(v, int):
        return format_brl(v)
    if col.kind == "bps" and isinstance(v, int):
        return f"{v / 100:.2f}%".replace(".", ",")
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    return v


def to_csv(ds: Dataset) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    w.writerow([c.label for c in ds.columns])
    for row in ds.rows:
        out = []
        for c in ds.columns:
            v = row.get(c.key)
            if c.kind == "money" and isinstance(v, int):
                v = f"{v / 100:.2f}".replace(".", ",")  # valor numérico em reais (centavos/100), sem símbolo
            else:
                v = _display(c, v)
            out.append(_safe_text(v) if isinstance(v, str) else v)
        w.writerow(out)
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(ds: Dataset, *, brand=None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = (ds.title or "Relatório")[:31].replace("/", "-")
    color = ((getattr(brand, "primary_color", None) or "#1D4ED8").lstrip("#")).upper()
    r = 1
    # cabeçalho com a marca da INSTÂNCIA (fallback RD); a logo vai nos PDFs (planilha mantém só nome/cor)
    ws.cell(row=r, column=1, value=_safe_text(getattr(brand, "name", None) or "RD Mídia Indoor")).font = Font(
        bold=True, size=11, color=color)
    r += 1
    ws.cell(row=r, column=1, value=_safe_text(ds.title)).font = Font(bold=True, size=13)
    r += 1
    for k, v in ds.meta:
        ws.cell(row=r, column=1, value=_safe_text(str(k))).font = Font(bold=True)
        ws.cell(row=r, column=2, value=_safe_text(str(v)))
        r += 1
    r += 1
    header_row = r
    fill = PatternFill("solid", fgColor=color)
    for i, c in enumerate(ds.columns, start=1):
        cell = ws.cell(row=r, column=i, value=c.label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    for row in ds.rows:
        r += 1
        for i, c in enumerate(ds.columns, start=1):
            v = row.get(c.key)
            cell = ws.cell(row=r, column=i)
            if c.kind == "money" and isinstance(v, int):
                cell.value = v / 100  # exibição em reais; o dado de origem é inteiro em centavos
                cell.number_format = '"R$" #,##0.00'
            elif c.kind == "bps" and isinstance(v, int):
                cell.value = v / 10000
                cell.number_format = "0.00%"
            elif c.kind == "int" and isinstance(v, int):
                cell.value = v
            elif isinstance(v, datetime):
                cell.value = v.replace(tzinfo=None)
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif isinstance(v, date):
                cell.value = v
                cell.number_format = "dd/mm/yyyy"
            elif v is None:
                cell.value = None
            else:
                cell.value = _safe_text(str(v))
    for note in ds.notes:
        r += 2
        ws.cell(row=r, column=1, value=_safe_text(note)).font = Font(italic=True, color="475569")
    for i, c in enumerate(ds.columns, start=1):
        width = max([len(c.label)] + [len(str(_display(c, row.get(c.key)) or "")) for row in ds.rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(ds: Dataset, *, brand=None, subtitle: Optional[str] = None, footer: Optional[str] = None) -> bytes:
    from app.modules.reports.pdf import KeyValues, Para, PdfDoc, Table, render_pdf

    numeric = [i for i, c in enumerate(ds.columns) if c.kind in ("int", "money", "bps")]
    rows = [[_display(c, row.get(c.key)) for c in ds.columns] for row in ds.rows]
    blocks: list = []
    if ds.meta:
        blocks.append(KeyValues(ds.meta))
    blocks.append(Table([c.label for c in ds.columns], rows, align_right=numeric))
    for n in ds.notes:
        blocks.append(Para(n, muted=True))
    bk = brand.pdf_kwargs() if brand is not None else {}
    return render_pdf(PdfDoc(title=ds.title, subtitle=subtitle, blocks=blocks, footer=footer,
                             landscape=len(ds.columns) > 6, **bk))


def export(ds: Dataset, fmt: str, *, brand=None, footer: Optional[str] = None) -> tuple[bytes, str, str]:
    """→ (bytes, content_type, extensão). ``brand`` = ``instance.service.Brand`` (marca da instância)."""
    fmt = fmt.lower()
    if fmt == "csv":
        return to_csv(ds), CONTENT_TYPES["csv"], "csv"
    if fmt == "xlsx":
        return to_xlsx(ds, brand=brand), CONTENT_TYPES["xlsx"], "xlsx"
    if fmt == "pdf":
        return to_pdf(ds, brand=brand, footer=footer), CONTENT_TYPES["pdf"], "pdf"
    raise ValueError("formato inválido")


def columns(*spec: Sequence[str]) -> list[Column]:
    return [Column(*s) for s in spec]
