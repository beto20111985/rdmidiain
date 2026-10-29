"""W-BIZ — Instância / white-label: piso de preço RD (preço abaixo do mínimo rejeitado),
branding restrito aos tokens do DS (contraste), referência ao segredo CloudPayPag sem
expor valor, logo, chave de assinatura ES256 da instância e permissões."""
from __future__ import annotations

import io
import os

import pytest
from PIL import Image
from sqlalchemy import select, text
from sqlalchemy.exc import DBAPIError

from app.core.errors import DomainError
from app.core.storage import LocalStorageProvider, set_storage
from app.models.instance import RdPriceFloorHistory
from app.modules.instance import service as inst
from tests.factories import login_as, make_user


@pytest.fixture(autouse=True)
async def _drain_after_test(clean_state):
    yield
    from app.modules.notifications.service import drain

    await drain()


@pytest.fixture
def storage(tmp_path):
    set_storage(LocalStorageProvider(tmp_path / "storage"))
    yield
    set_storage(None)


async def test_price_below_rd_minimum_rejected(client, db, mail):
    master = await make_user(db, role="super_admin", mfa=True, admin_master_seat=1)
    other_sa = await make_user(db, role="super_admin", mfa=True)
    sm, so = await login_as(client, master), await login_as(client, other_sa)
    body = {"label": "Anúncio 30 s — mensal por tela", "unit": "tela/mês", "min_price_cents": 15000,
            "reason": "Tabela base RD 2026"}
    r = await client.put("/api/v1/instance/price-floors/slot_30s_mensal_tela", headers=so.headers, json=body)
    assert r.status_code == 403 and r.json()["error"]["code"] == "ADMIN_MASTER_REQUIRED"
    r = await client.put("/api/v1/instance/price-floors/slot_30s_mensal_tela", headers=sm.headers, json=body)
    assert r.status_code == 200, r.text
    # API de verificação (usada pela UI de preços) e serviço (usado pelo W-CAMP)
    r = await client.post("/api/v1/instance/price-floors/check", headers=sm.headers,
                          json={"floor_key": "slot_30s_mensal_tela", "price_cents": 14999})
    assert r.status_code == 400 and r.json()["error"]["code"] == "PRICE_BELOW_RD_MINIMUM"
    assert r.json()["error"]["details"]["min_price_cents"] == 15000
    r = await client.post("/api/v1/instance/price-floors/check", headers=sm.headers,
                          json={"floor_key": "slot_30s_mensal_tela", "price_cents": 15000})
    assert r.status_code == 200 and r.json()["ok"] is True
    with pytest.raises(DomainError) as exc:
        await inst.assert_price_not_below_floor(db, "slot_30s_mensal_tela", 100)
    assert exc.value.code == "PRICE_BELOW_RD_MINIMUM"
    await inst.assert_price_not_below_floor(db, "slot_30s_mensal_tela", 20000)
    with pytest.raises(TypeError):
        await inst.assert_price_not_below_floor(db, "slot_30s_mensal_tela", 149.99)  # float nunca
    # histórico do piso, append-only
    hist = (await db.execute(select(RdPriceFloorHistory))).scalars().all()
    assert len(hist) == 1 and hist[0].new_min_price_cents == 15000
    r = await client.put("/api/v1/instance/price-floors/slot_30s_mensal_tela", headers=sm.headers,
                         json={**body, "min_price_cents": 17000, "reason": "Reajuste anual", "expected_version": 1})
    assert r.status_code == 200 and r.json()["version"] == 2
    h = await client.get("/api/v1/instance/price-floors/slot_30s_mensal_tela/history", headers=sm.headers)
    assert h.status_code == 200, h.text
    items = h.json()["items"]
    assert h.json()["total"] == 2 and [i["version"] for i in items] == [2, 1]
    assert items[0]["old_min_price_cents"] == 15000 and items[0]["new_min_price_cents"] == 17000
    assert items[0]["reason"] == "Reajuste anual" and items[0]["changed_by"] == str(master.id)
    admin = await make_user(db, role="admin", mfa=True)  # audit.view, sem whitelabel.manage
    assert (await client.get("/api/v1/instance/price-floors/slot_30s_mensal_tela/history",
                             headers=(await login_as(client, admin)).headers)).status_code == 200
    sales = await make_user(db, role="sales")  # pricing.view, sem audit.view/whitelabel.manage
    assert (await client.get("/api/v1/instance/price-floors/slot_30s_mensal_tela/history",
                             headers=(await login_as(client, sales)).headers)).status_code == 403
    assert (await client.get("/api/v1/instance/price-floors/nao_existe/history", headers=sm.headers)).status_code == 404
    with pytest.raises(DBAPIError):
        await db.execute(text("UPDATE rd_price_floor_history SET new_min_price_cents = 1"))
    await db.rollback()


async def test_branding_tokens_and_secret_reference(client, db, storage, mail):
    master = await make_user(db, role="super_admin", mfa=True, admin_master_seat=2)
    s = await login_as(client, master)
    r = await client.patch("/api/v1/instance/profile", headers=s.headers, json={"primary_color": "#FFFF00"})
    assert r.status_code == 400 and r.json()["error"]["code"] == "BRAND_COLOR_CONTRAST"
    r = await client.patch("/api/v1/instance/profile", headers=s.headers, json={"accent_color": "#101010"})
    assert r.status_code == 400 and r.json()["error"]["code"] == "BRAND_COLOR_CONTRAST"
    r = await client.patch("/api/v1/instance/profile", headers=s.headers,
                           json={"payment_credentials_ref": "cpg_live_9f8e7d6c5b4a"})
    assert r.status_code == 400 and r.json()["error"]["code"] == "SECRET_REF_INVALID"
    os.environ["RD_TEST_CPP_TOKEN"] = "segredo-de-teste-nao-exibir"
    try:
        r = await client.patch("/api/v1/instance/profile", headers=s.headers, json={
            "display_name": "Mídia Franquia GO", "is_white_label": True, "primary_color": "#0056B3",
            "accent_color": "#5BCEFF", "subdomain": "franquia-go", "payment_credentials_ref": "env:RD_TEST_CPP_TOKEN"})
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["payment_credentials"]["configured"] is True
        assert "segredo-de-teste-nao-exibir" not in r.text
        assert body["isolation"]["model"] == "STACK_POR_INSTANCIA"
        assert await inst.resolve_payment_secret(db) == "segredo-de-teste-nao-exibir"  # uso interno do W-FIN
    finally:
        os.environ.pop("RD_TEST_CPP_TOKEN", None)
    pub = await client.get("/api/v1/instance/branding")
    assert pub.status_code == 200 and pub.json()["display_name"] == "Mídia Franquia GO"
    assert "payment_credentials_ref" not in pub.text and "env:" not in pub.text
    # logo: PNG aceito, SVG recusado; logo pública
    b = io.BytesIO()
    Image.new("RGB", (64, 64), (0, 86, 179)).save(b, "PNG")
    r = await client.put("/api/v1/instance/logo", headers=s.headers, files={"file": ("logo.png", b.getvalue(), "image/png")})
    assert r.status_code == 200, r.text
    r = await client.put("/api/v1/instance/logo", headers=s.headers,
                         files={"file": ("x.svg", b"<svg onload='alert(1)'></svg>", "image/svg+xml")})
    assert r.status_code == 415
    g = await client.get("/api/v1/instance/logo")
    assert g.status_code == 200 and g.content.startswith(b"\x89PNG")
    # admin (sem whitelabel.manage) não altera a marca
    admin = await make_user(db, role="admin", mfa=True)
    sa = await login_as(client, admin)
    assert (await client.patch("/api/v1/instance/profile", headers=sa.headers,
                               json={"display_name": "Hack"})).status_code == 403


async def test_instance_signing_key_es256(db):
    kid, sig = await inst.sign_payload(db, b"conteudo")
    await db.commit()
    assert await inst.verify_signature(db, kid, b"conteudo", sig) is True
    assert await inst.verify_signature(db, kid, b"conteudo alterado", sig) is False
    kid2, _ = await inst.sign_payload(db, b"outro")
    assert kid2 == kid  # uma única chave ativa
    pem = (await inst.public_keys(db))[0]["public_key_pem"]
    assert "BEGIN PUBLIC KEY" in pem and "PRIVATE" not in pem
    row = (await db.execute(text("SELECT private_enc FROM instance_signing_keys"))).scalar_one()
    assert row.startswith("v1:") and "PRIVATE KEY" not in row  # privada cifrada em repouso


async def test_white_label_brand_applied_to_documents_emails_and_labels(client, db, storage, mail):
    """MAT-1 / REQ-NEG-005: PDF, XLSX, e-mail e templates usam a marca da INSTÂNCIA (sem domínio fixo)."""
    import io as _io

    from openpyxl import load_workbook
    from pypdf import PdfReader

    from app.core.config import settings
    from app.modules.notifications import service as notif
    from app.modules.support.people import describe_authors

    master = await make_user(db, role="super_admin", mfa=True, admin_master_seat=1)
    s = await login_as(client, master)
    # padrão (sem configuração): fallback RD
    pub = await client.get("/api/v1/instance/branding")
    assert pub.json()["display_name"] == "RD Mídia Indoor"
    r = await client.patch("/api/v1/instance/profile", headers=s.headers, json={
        "display_name": "Mídia Franquia GO", "is_white_label": True, "primary_color": "#0B5D1E",
        "accent_color": "#5BCEFF"})
    assert r.status_code == 200, r.text
    b = io.BytesIO()
    Image.new("RGB", (64, 64), (200, 30, 30)).save(b, "PNG")
    assert (await client.put("/api/v1/instance/logo", headers=s.headers,
                             files={"file": ("logo.png", b.getvalue(), "image/png")})).status_code == 200

    # /branding: só campos seguros, ETag e revalidação (304)
    pub = await client.get("/api/v1/instance/branding")
    assert pub.status_code == 200
    assert set(pub.json()) == {"instance_slug", "display_name", "is_white_label", "primary_color", "accent_color",
                               "has_logo", "support_email", "support_phone", "logo_url"}
    assert pub.json()["logo_url"].startswith("/api/v1/instance/logo?v=")
    etag = pub.headers["etag"]
    assert "max-age" in pub.headers["cache-control"]
    assert (await client.get("/api/v1/instance/branding", headers={"If-None-Match": etag})).status_code == 304
    lg = await client.get("/api/v1/instance/logo")
    assert (await client.get("/api/v1/instance/logo", headers={"If-None-Match": lg.headers["etag"]})).status_code == 304

    # PDF e XLSX exportados: nome, cor e logo da instância (nada de "RD Mídia Indoor")
    files = {}
    for fmt in ("pdf", "xlsx"):
        e = (await client.post("/api/v1/reports/exports", headers=s.headers,
                               json={"report_key": "exibicoes_por_dia", "format": fmt})).json()
        files[fmt] = (await client.get(e["download_url"], headers=s.headers)).content
    pdf = PdfReader(_io.BytesIO(files["pdf"]))
    txt = "".join(p.extract_text() for p in pdf.pages)
    assert "Mídia Franquia GO" in txt and "RD Mídia Indoor" not in txt
    assert len(pdf.pages[0].images) >= 1  # logo desenhada no cabeçalho
    ws = load_workbook(_io.BytesIO(files["xlsx"])).active
    assert ws["A1"].value == "Mídia Franquia GO"
    header = next(r for r in ws.iter_rows() if r[0].value == "Dia")
    assert header[0].fill.fgColor.rgb.endswith("0B5D1E")

    # e-mail de notificação: assunto/rodapé/HTML na marca; links da configuração; template com {{instance_name}}
    ok = await client.put("/api/v1/notifications/templates/campaign.approved/EMAIL", headers=s.headers, json={
        "title_template": "{{instance_name}}: {{campaign_name}} aprovada", "body_template": "Tudo certo."})
    assert ok.status_code == 200, ok.text
    u = await make_user(db, name="Cliente")
    await notif.notify(db, [u.id], "campaign.approved", {"campaign_name": "Inverno", "campaign_id": "9"})
    await notif.notify(db, [u.id], "campaign.awaiting_approval", {"campaign_name": "Verão", "campaign_id": "8"},
                       channels=["EMAIL"])
    await db.commit()
    await notif.drain()
    msgs = [m for m in mail.outbox if m.to == u.email]
    approved = next(m for m in msgs if "aprovada" in m.subject)
    assert approved.subject == "Mídia Franquia GO: Inverno aprovada"
    assert "Mídia Franquia GO." in approved.text and "RD Mídia Indoor" not in approved.text
    base = settings.panel_base_url
    assert f"{base}/notificacoes/preferencias" in approved.text
    assert approved.html and "Mídia Franquia GO" in approved.html and "#0B5D1E" in approved.html
    assert f"{base}/api/v1/instance/logo" in approved.html and "rdmidiaindoor" not in approved.html
    awaiting = next(m for m in msgs if m is not approved)
    assert "revisão da equipe Mídia Franquia GO" in awaiting.text
    # rótulo da equipe exibido a clientes acompanha a instância
    staff = await make_user(db, role="employee", name="Ana Suporte")
    names = await describe_authors(db, {staff.id}, viewer_is_staff=False, viewer_user_id=u.id)
    assert names[staff.id]["author_name"] == "Equipe Mídia Franquia GO"
