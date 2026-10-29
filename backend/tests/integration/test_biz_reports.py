"""W-BIZ — Relatórios, exportações reais (CSV/XLSX/PDF abertos de verdade), e-mail com
anexo, comprovante de veiculação ASSINADO contendo só PoP válido (e sem nomes de
estabelecimentos para o anunciante), relatório mensal automático e dashboards com
números conferidos contra os dados de teste e escopo respeitado.

Dados de PoP: eventos brutos gravados como o W-PLAY grava (``playback_events``) e
agregados pela rotina REAL do W-PLAY (``aggregate_pending``)."""
from __future__ import annotations

import base64
import csv
import hashlib
import io
import secrets
import uuid
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import delete, select, text, update
from sqlalchemy.exc import DBAPIError

from app.core.storage import LocalStorageProvider, set_storage
from app.core.timeutils import local_date, utcnow
from app.models.reports import DeliveryProof, ReportRun
from app.modules.reports import data as rdata
from app.modules.reports import service as rsvc
from tests.factories import login_as, make_user
from tests.inv_support import link, mk_campaign, mk_creative, mk_establishment, mk_org, mk_screen


@pytest.fixture(autouse=True)
async def _drain_after_test(clean_state):
    yield
    from app.modules.notifications.service import drain

    await drain()


@pytest.fixture
def storage(tmp_path):
    set_storage(LocalStorageProvider(tmp_path / "storage"))
    yield
    set_storage(None)


async def _device(db, screen, *, online: bool):
    from app.models.devices import Device
    from app.models.screens import Screen

    d = Device(id=uuid.uuid4(), screen_id=screen.id, public_key="x", key_alg="ES256",
               key_fingerprint=secrets.token_hex(32), platform="web", status="ACTIVE", paired_at=utcnow(),
               pairing_method="PANEL", presence="online" if online else "offline",
               last_seen_at=utcnow() - (timedelta(minutes=1) if online else timedelta(days=9)),
               applied_manifest_version=3, pending_manifest_version=None if online else 4)
    db.add(d)
    await db.flush()
    await db.execute(update(Screen).where(Screen.id == screen.id).values(device_id=d.id))
    await db.commit()
    return d


_seq = 0
INTERNAL_CREATIVE = uuid.uuid4()  # uma única peça interna do estabelecimento


def _ev(dev, screen, est, *, creative, campaign, kind="EXTERNAL", when, valid=True, status="ACCEPTED",
        reason=None, ms=30000):
    from app.models.playback import PlaybackEvent

    global _seq
    _seq += 1
    st_ms = int(when.timestamp() * 1000)
    return PlaybackEvent(
        playback_uuid=uuid.uuid4(), device_id=dev.id, screen_id=screen.id, establishment_id=est.id, org_id=est.org_id,
        creative_id=creative.id if creative else INTERNAL_CREATIVE, creative_code=creative.code if creative else "RD-INTERN",
        campaign_id=campaign.id if campaign else None, kind=kind, manifest_version=3, started_at_ms=st_ms,
        finished_at_ms=st_ms + ms, started_at=when, finished_at=when + timedelta(milliseconds=ms), expected_ms=30000,
        actual_ms=ms, result="COMPLETED" if valid else "INTERRUPTED", seq=_seq, nonce=secrets.token_hex(8),
        prev_hash="0" * 64, event_hash=secrets.token_hex(32), clock_offset_ms=0, offline=False, signature="x",
        status=status, valid=valid, reason=reason, reasons=[] if valid else [reason or "X"], flags=[],
        local_day=local_date(when, est.timezone))


async def _world(db, client):
    adv_a, adv_b = await mk_org(db, "Anunciante Alfa"), await mk_org(db, "Anunciante Beta")
    est_org = await mk_org(db, "Supermercado Bom Preço")
    est = await mk_establishment(db, est_org)  # Goiânia/GO
    s1, s2 = await mk_screen(db, est), await mk_screen(db, est)
    d1, d2 = await _device(db, s1, online=True), await _device(db, s2, online=False)
    cr_a, cr_b = await mk_creative(db, adv_a), await mk_creative(db, adv_b)
    ca = await mk_campaign(db, adv_a, status="ATIVA", name="Campanha Alfa")
    cb = await mk_campaign(db, adv_b, status="EM_ANALISE", name="Campanha Beta")
    await link(db, ca, cr_a)
    await link(db, cb, cr_b)
    now = utcnow() - timedelta(minutes=30)
    old = utcnow() - timedelta(days=120)
    evs = []
    evs += [_ev(d1, s1, est, creative=cr_a, campaign=ca, when=now - timedelta(minutes=i)) for i in range(5)]
    evs += [_ev(d1, s1, est, creative=cr_a, campaign=ca, when=now - timedelta(minutes=10 + i), valid=False,
                reason="INCOMPLETE_PLAYBACK", ms=25000) for i in range(2)]
    evs += [_ev(d1, s1, est, creative=cr_a, campaign=ca, when=now - timedelta(minutes=20), valid=False,
                status="REJECTED", reason="BAD_SIGNATURE")]
    evs += [_ev(d2, s2, est, creative=cr_a, campaign=ca, when=now - timedelta(minutes=i)) for i in range(3)]
    evs += [_ev(d1, s1, est, creative=cr_b, campaign=cb, when=now - timedelta(minutes=30 + i)) for i in range(4)]
    evs += [_ev(d1, s1, est, creative=None, campaign=None, kind="INTERNAL", when=now - timedelta(minutes=40 + i))
            for i in range(6)]
    evs += [_ev(d2, s2, est, creative=cr_a, campaign=ca, when=old - timedelta(minutes=i)) for i in range(7)]
    for e in evs:
        db.add(e)
    await db.commit()
    from app.modules.playback.maintenance import aggregate_pending

    await aggregate_pending(db)
    # retenção: brutos com mais de 90 dias já foram expurgados (agregado preservado)
    from app.models.playback import PlaybackEvent

    await db.execute(delete(PlaybackEvent).where(PlaybackEvent.started_at < utcnow() - timedelta(days=90)))
    from app.models.devices import DeviceTelemetry

    for i in range(12):  # 1 hora de amostras (5 min) para a tela 1
        db.add(DeviceTelemetry(device_id=d1.id, screen_id=s1.id, recorded_at=utcnow() - timedelta(minutes=5 * i),
                               data={"cpu_pct": 10}))
    await db.commit()
    users = {
        "adv_a": await make_user(db, role="advertiser", scope=("ORGANIZATION", adv_a.id)),
        "adv_b": await make_user(db, role="advertiser", scope=("ORGANIZATION", adv_b.id)),
        "owner": await make_user(db, role="establishment_owner", scope=("ORGANIZATION", est_org.id)),
        "admin": await make_user(db, role="admin", mfa=True),
    }
    sess = {k: await login_as(client, u) for k, u in users.items()}
    return dict(adv_a=adv_a, adv_b=adv_b, est_org=est_org, est=est, s1=s1, s2=s2, ca=ca, cb=cb, cr_a=cr_a, cr_b=cr_b,
                old=old, users=users, s=sess)


def _period_params(w, days_back=0):
    today = local_date(utcnow(), "America/Sao_Paulo")
    start = today - timedelta(days=days_back) if days_back else today
    return {"start": start.isoformat(), "end": today.isoformat()}


async def test_report_scopes_advertiser_establishment_staff(client, db, storage, mail):
    w = await _world(db, client)
    sa, so, sadm = w["s"]["adv_a"], w["s"]["owner"], w["s"]["admin"]
    q = _period_params(w)
    r = await client.get("/api/v1/reports/data/exibicoes_por_campanha", headers=sa.headers,
                         params={**q, "audience": "ANUNCIANTE", "org_id": str(w["adv_a"].id)})
    assert r.status_code == 200, r.text
    rows = r.json()["rows"]
    assert [x["campaign"] for x in rows] == ["Campanha Alfa"] and rows[0]["valid"] == 8 and rows[0]["invalid"] == 3
    assert "advertiser" not in rows[0] or rows[0]["advertiser"] is None
    r = await client.get("/api/v1/reports/data/exibicoes_por_tela", headers=sa.headers,
                         params={**q, "audience": "ANUNCIANTE", "org_id": str(w["adv_a"].id)})
    body = r.json()
    assert {c["key"] for c in body["columns"]} == {"screen", "region", "valid", "valid_hours"}
    assert "Supermercado" not in r.text and "Goiânia" not in r.text and "Est " not in r.text
    assert all(row["region"].startswith("GO") for row in body["rows"])
    # anunciante não acessa relatórios de estabelecimento/falhas nem dados de outra org
    for key in ("exibicoes_por_estabelecimento", "falhas", "uptime"):
        r = await client.get(f"/api/v1/reports/data/{key}", headers=sa.headers,
                             params={**q, "audience": "ANUNCIANTE", "org_id": str(w["adv_a"].id)})
        assert r.status_code == 403, key
    r = await client.get("/api/v1/reports/data/exibicoes_por_campanha", headers=sa.headers,
                         params={**q, "audience": "ANUNCIANTE", "org_id": str(w["adv_b"].id)})
    assert r.status_code == 404
    # estabelecimento: por tipo, sem anunciante/campanha/ID do anúncio
    r = await client.get("/api/v1/reports/data/exibicoes_por_tipo", headers=so.headers,
                         params={**q, "audience": "ESTABELECIMENTO", "org_id": str(w["est_org"].id)})
    assert r.status_code == 200, r.text
    by = {x["kind"]: x["valid"] for x in r.json()["rows"]}
    assert by == {"Externo (anúncios)": 12, "Interno (estabelecimento)": 6}
    assert "Alfa" not in r.text and "Beta" not in r.text and w["cr_a"].code not in r.text
    for key in ("exibicoes_por_campanha", "exibicoes_por_anuncio"):
        r = await client.get(f"/api/v1/reports/data/{key}", headers=so.headers,
                             params={**q, "audience": "ESTABELECIMENTO", "org_id": str(w["est_org"].id)})
        assert r.status_code == 403
    r = await client.get("/api/v1/reports/data/exibicoes_por_horario", headers=so.headers,
                         params={**q, "audience": "ESTABELECIMENTO", "org_id": str(w["est_org"].id)})
    assert sum(x["valid"] for x in r.json()["rows"]) == 18
    r = await client.get("/api/v1/reports/data/uptime", headers=so.headers,
                         params={**q, "audience": "ESTABELECIMENTO", "org_id": str(w["est_org"].id)})
    up = {x["screen"]: x for x in r.json()["rows"]}
    assert up[w["s1"].friendly_id]["uptime_bps"] > 0 and up[w["s2"].friendly_id]["uptime_bps"] == 0
    # equipe: falhas por motivo e por estabelecimento
    r = await client.get("/api/v1/reports/data/falhas", headers=sadm.headers, params=q)
    reasons = {x["reason"]: x["count"] for x in r.json()["rows"]}
    assert reasons == {"INCOMPLETE_PLAYBACK": 2, "BAD_SIGNATURE": 1}
    r = await client.get("/api/v1/reports/data/exibicoes_por_estabelecimento", headers=sadm.headers,
                         params=_period_params(w, 150))
    assert r.json()["rows"][0]["valid"] == 25  # 18 recentes + 7 antigas (agregado preservado)


async def test_exports_csv_xlsx_pdf_are_real_and_email_has_attachment(client, db, storage, mail):
    w = await _world(db, client)
    sadm = w["s"]["admin"]
    files = {}
    for fmt in ("csv", "xlsx", "pdf"):
        r = await client.post("/api/v1/reports/exports", headers=sadm.headers, json={
            "report_key": "exibicoes_por_tela", "format": fmt, **_period_params(w, 150), "email_to_me": fmt == "xlsx"})
        assert r.status_code == 201, r.text
        e = r.json()
        dl = await client.get(e["download_url"], headers=sadm.headers)
        assert dl.status_code == 200
        assert hashlib.sha256(dl.content).hexdigest() == e["sha256"]
        files[fmt] = dl.content
    rows = list(csv.reader(io.StringIO(files["csv"].decode("utf-8-sig")), delimiter=";"))
    assert rows[0][0] == "Tela (código)"
    by_screen = {r[0]: r for r in rows[1:]}
    assert by_screen[w["s2"].friendly_id][4] == "10"  # 3 recentes + 7 antigas
    wb = load_workbook(io.BytesIO(files["xlsx"]))
    ws = wb.active
    values = [[c.value for c in row] for row in ws.iter_rows()]
    header_idx = next(i for i, r in enumerate(values) if r[0] == "Tela (código)")
    data = {r[0]: r for r in values[header_idx + 1:] if r[0]}
    assert data[w["s1"].friendly_id][4] == 15  # 5 + 4 externas + 6 internas
    txt = "".join(p.extract_text() for p in PdfReader(io.BytesIO(files["pdf"])).pages)
    assert "Exibições por tela" in txt and w["s1"].friendly_id in txt
    m = mail.last_to(w["users"]["admin"].email)
    assert m is not None and len(m.attachments) == 1 and m.attachments[0][0].endswith(".xlsx")
    assert m.attachments[0][2] == files["xlsx"]
    # download só pelo solicitante
    other = await login_as(client, w["users"]["adv_a"])
    assert (await client.get(f"/api/v1/reports/exports/{e['id']}/download", headers=other.headers)).status_code == 404


async def test_delivery_proof_signed_only_valid_pop_no_establishment_names(client, db, storage, mail):
    w = await _world(db, client)
    sa = w["s"]["adv_a"]
    start = local_date(w["old"], "America/Sao_Paulo") - timedelta(days=1)
    r = await client.post("/api/v1/reports/proofs", headers=sa.headers, json={
        "campaign_id": str(w["ca"].id), "start": start.isoformat(),
        "end": local_date(utcnow(), "America/Sao_Paulo").isoformat()})
    assert r.status_code == 201, r.text
    proof = r.json()
    # 5 + 3 válidas recentes (brutos) + 7 antigas (agregado de válidas); NÃO entram 2 inválidas + 1 rejeitada,
    # nem a campanha Beta nem o conteúdo interno
    assert proof["total_valid_plays"] == 15
    assert proof["source"] == "playback_events.valid+playback_daily_stats.plays_valid"
    pdf = (await client.get(proof["pdf_url"], headers=sa.headers)).content
    assert pdf.startswith(b"%PDF") and hashlib.sha256(pdf).hexdigest() == proof["pdf_sha256"]
    txt = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert proof["code"] in txt and w["cr_a"].code in txt and w["s1"].friendly_id in txt
    assert "Supermercado" not in txt and "Goiânia" not in txt and "GO (" in txt
    assert "Página de verificação" in txt and "ICP-Brasil" in txt
    # verificação pública pelo PDF
    v = await client.post("/api/v1/reports/proofs/verify", files={"file": ("c.pdf", pdf, "application/pdf")})
    assert v.status_code == 200 and v.json()["valid"] is True and v.json()["code"] == proof["code"]
    tampered = bytearray(pdf)
    tampered[len(tampered) // 2] ^= 0x01
    v = await client.post("/api/v1/reports/proofs/verify", files={"file": ("c.pdf", bytes(tampered), "application/pdf")})
    assert v.json()["valid"] is False
    # verificação OFFLINE: assinatura ES256 do payload canônico com a chave pública publicada
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import ec

    row = (await db.execute(select(DeliveryProof).where(DeliveryProof.code == proof["code"]))).scalar_one()
    keys = (await client.get("/api/v1/instance/signing-keys")).json()["items"]
    pub = serialization.load_pem_public_key(next(k for k in keys if k["kid"] == row.kid)["public_key_pem"].encode())
    sig = base64.urlsafe_b64decode(row.payload_signature + "=" * (-len(row.payload_signature) % 4))
    canon = rsvc.canonical(row.payload)
    pub.verify(sig, canon, ec.ECDSA(hashes.SHA256()))  # lança se inválida
    assert hashlib.sha256(canon).hexdigest() == row.payload_sha256
    assert all(s["valid_plays"] > 0 for s in row.payload["screens"])
    assert {s["creative_code"] for s in row.payload["screens"]} == {w["cr_a"].code}
    # comprovante é imutável; outra organização não emite comprovante da campanha alheia
    with pytest.raises(DBAPIError):
        await db.execute(text("UPDATE delivery_proofs SET total_valid_plays = 999"))
    await db.rollback()
    r = await client.post("/api/v1/reports/proofs", headers=w["s"]["adv_b"].headers, json={"campaign_id": str(w["ca"].id)})
    assert r.status_code == 404
    # campanha sem PoP: zero (agenda planejada nunca vira execução)
    cz = await mk_campaign(db, w["adv_a"], status="ATIVA", name="Sem exibição")
    r = await client.post("/api/v1/reports/proofs", headers=sa.headers, json={"campaign_id": str(cz.id)})
    assert r.status_code == 201 and r.json()["total_valid_plays"] == 0


async def test_monthly_report_emailed_once(client, db, storage, mail):
    w = await _world(db, client)
    today = local_date(utcnow(), "America/Sao_Paulo")
    period = rdata.Period(today - timedelta(days=2), today)
    stats = await rsvc.run_monthly(db, period)
    await db.commit()
    assert stats["advertisers"] == 2 and stats["establishments"] == 1 and stats["emails"] == 3
    for k in ("adv_a", "adv_b", "owner"):
        m = mail.last_to(w["users"][k].email)
        assert m is not None and m.attachments and m.attachments[0][2].startswith(b"%PDF"), k
    txt_a = "".join(p.extract_text() for p in PdfReader(io.BytesIO(mail.last_to(w["users"]["adv_a"].email).attachments[0][2])).pages)
    assert "Campanha Alfa" in txt_a and "Campanha Beta" not in txt_a and "Supermercado" not in txt_a
    txt_o = "".join(p.extract_text() for p in PdfReader(io.BytesIO(mail.last_to(w["users"]["owner"].email).attachments[0][2])).pages)
    assert "Alfa" not in txt_o and "Beta" not in txt_o
    again = await rsvc.run_monthly(db, period)
    await db.commit()
    assert again["skipped_existing"] == 3 and again["emails"] == 0
    assert len((await db.execute(select(ReportRun))).scalars().all()) == 3


async def test_dashboards_numbers_and_scope(client, db, storage, mail):
    w = await _world(db, client)
    r = await client.get("/api/v1/dashboards/admin", headers=w["s"]["admin"].headers)
    assert r.status_code == 200, r.text
    k = r.json()["kpis"]
    assert k["screens_total"] == 2 and k["screens_online"] == 1 and k["players_offline"] == 1
    assert k["campaigns_active"] == 1 and k["advertisers_active"] == 1
    assert k["exhibitions_today"] == 18
    ops = r.json()["operations"]
    assert ops["campaigns_awaiting_approval"] == 1 and ops["screens_outdated"] == 1 and ops["screens_offline_long"] == 1
    assert r.json()["distribution_by_state"] == [{"uf": "GO", "screens": 2, "online": 1, "exhibitions_period": 18}]
    # estabelecimento: seus números, sem anunciante e sem saque
    r = await client.get(f"/api/v1/dashboards/establishment?org_id={w['est_org'].id}", headers=w["s"]["owner"].headers)
    assert r.status_code == 200, r.text
    e = r.json()
    assert e["kpis"]["screens_total"] == 2 and e["kpis"]["screens_online"] == 1
    assert e["kpis"]["external_exhibitions_month"] == 12 and e["kpis"]["internal_exhibitions_month"] == 6
    assert e["programming"]["external_ads_running_today"] == 2 and e["programming"]["internal_contents_today"] == 1
    assert "Alfa" not in r.text and "Beta" not in r.text and w["cr_a"].code not in r.text
    assert "withdrawals" in e["commission"] and isinstance(e["commission"]["withdrawals"], str)
    # anunciante: somente as próprias campanhas, regiões sem nomes de estabelecimento
    r = await client.get(f"/api/v1/dashboards/advertiser?org_id={w['adv_a'].id}", headers=w["s"]["adv_a"].headers)
    assert r.status_code == 200, r.text
    a = r.json()
    assert [c["name"] for c in a["campaigns"]] == ["Campanha Alfa"]
    assert a["regions_exhibited"] == [{"uf": "GO", "screens": 2}]
    assert "Supermercado" not in r.text and "Beta" not in r.text
    # cross-tenant e perfis errados
    assert (await client.get(f"/api/v1/dashboards/advertiser?org_id={w['adv_b'].id}",
                             headers=w["s"]["adv_a"].headers)).status_code == 404
    assert (await client.get(f"/api/v1/dashboards/establishment?org_id={w['est_org'].id}",
                             headers=w["s"]["adv_a"].headers)).status_code == 404
    assert (await client.get("/api/v1/dashboards/admin", headers=w["s"]["owner"].headers)).status_code == 403
    # capacidade DISPONÍVEL sem ocupação (E2#717 NÃO / #718 SIM) — falha se a ocupação voltar ao payload
    cap = await client.get("/api/v1/dashboards/admin/capacity", headers=w["s"]["admin"].headers)
    assert cap.status_code == 200, cap.text
    c = cap.json()
    assert set(c) == {"date", "screens", "free_units", "unit", "unit_seconds", "free_seconds", "truncated", "note"}
    assert c["unit"] == "SLOT_30S" and c["unit_seconds"] == 30 and c["free_seconds"] == c["free_units"] * 30
    assert "occupancy" not in cap.text and "used_units" not in cap.text and "capacity_units" not in cap.text
    # texto para o usuário: sem "ocupação" e sem referência interna a fontes (ETAPA/#)
    assert "ocupa" not in cap.text.lower() and "etapa" not in cap.text.lower() and "#7" not in cap.text
    adm = (await client.get("/api/v1/dashboards/admin", headers=w["s"]["admin"].headers)).text
    assert "occupancy" not in adm and "ocupa" not in adm.lower() and "etapa" not in adm.lower()


async def test_downloads_and_sync_reports_from_player_history(client, db, storage, mail):
    from app.models.devices import Device
    from app.modules.playback import syncevents

    w = await _world(db, client)
    other_org = await mk_org(db, "Academia Outra")
    other_est = await mk_establishment(db, other_org)
    other_screen = await mk_screen(db, other_est)
    other_dev = await _device(db, other_screen, online=True)
    d1 = (await db.execute(select(Device).where(Device.screen_id == w["s1"].id))).scalar_one()
    d2 = (await db.execute(select(Device).where(Device.screen_id == w["s2"].id))).scalar_one()
    sc_est = {"org_id": str(w["est"].org_id), "establishment_id": str(w["est"].id)}
    sc_other = {"org_id": str(other_est.org_id), "establishment_id": str(other_est.id)}
    mb = 1024 * 1024
    for ev in (
        syncevents.build(w["s1"].id, d1.id, "MANIFEST_RECEIVED", scope=sc_est, manifest_version=4),
        syncevents.build(w["s1"].id, d1.id, "DOWNLOAD_STARTED", scope=sc_est, manifest_version=4),
        syncevents.build(w["s1"].id, d1.id, "DOWNLOAD_COMPLETED", scope=sc_est, manifest_version=4, bytes_=5 * mb),
        syncevents.build(w["s1"].id, d1.id, "MANIFEST_APPLIED", scope=sc_est, manifest_version=4, duration_ms=4000),
        syncevents.build(w["s2"].id, d2.id, "MANIFEST_RECEIVED", scope=sc_est, manifest_version=4),
        syncevents.build(w["s2"].id, d2.id, "DOWNLOAD_STARTED", scope=sc_est, manifest_version=4),
        syncevents.build(w["s2"].id, d2.id, "MEDIA_ERROR", scope=sc_est, media_id=uuid.uuid4(), reason="CHECKSUM"),
        syncevents.build(w["s2"].id, d2.id, "MANIFEST_FAILED", scope=sc_est, manifest_version=4, reason="DOWNLOAD"),
        syncevents.build(other_screen.id, other_dev.id, "DOWNLOAD_COMPLETED", scope=sc_other, bytes_=1 * mb),
    ):
        db.add(ev)
    await db.commit()
    q = _period_params(w)
    sadm, so, sa = w["s"]["admin"], w["s"]["owner"], w["s"]["adv_a"]

    r = await client.get("/api/v1/reports/data/downloads", headers=sadm.headers, params=q)
    assert r.status_code == 200, r.text
    rows = {x["screen"]: x for x in r.json()["rows"]}
    assert set(rows) == {w["s1"].friendly_id, w["s2"].friendly_id, other_screen.friendly_id}
    s1 = rows[w["s1"].friendly_id]
    assert (s1["dl_started"], s1["dl_completed"], s1["volume_mb"], s1["received"], s1["applied"]) == (1, 1, "5,0", 1, 1)
    assert s1["p50_s"] == "4,0" and s1["establishment"] is not None
    s2 = rows[w["s2"].friendly_id]
    assert (s2["dl_completed"], s2["failed"], s2["media_errors"]) == (0, 1, 1)
    meta = {m["label"]: m["value"] for m in r.json()["meta"]}
    assert meta["Downloads concluídos"] == "2" and meta["Volume total baixado (MB)"] == "6,0"
    assert meta["Mídias distintas com erro"] == "1"

    # estabelecimento: só as suas telas, sem coluna de estabelecimento nem detalhe de mídia
    r = await client.get("/api/v1/reports/data/downloads", headers=so.headers,
                         params={**q, "audience": "ESTABELECIMENTO", "org_id": str(w["est_org"].id)})
    assert r.status_code == 200, r.text
    assert {x["screen"] for x in r.json()["rows"]} == {w["s1"].friendly_id, w["s2"].friendly_id}
    assert "establishment" not in {c["key"] for c in r.json()["columns"]}
    assert "Mídias distintas com erro" not in {m["label"] for m in r.json()["meta"]}
    # anunciante não tem acesso
    r = await client.get("/api/v1/reports/data/downloads", headers=sa.headers,
                         params={**q, "audience": "ANUNCIANTE", "org_id": str(w["adv_a"].id)})
    assert r.status_code == 403

    # sincronizações por dia (equipe): totais do dia
    r = await client.get("/api/v1/reports/data/sincronizacoes_por_dia", headers=sadm.headers, params=q)
    assert r.status_code == 200, r.text
    (day,) = r.json()["rows"]
    assert (day["received"], day["applied"], day["failed"], day["dl_completed"], day["volume_mb"], day["media_errors"]) \
        == (2, 1, 1, 2, "6,0", 1)
    # exportação real (CSV) do relatório de downloads pelo estabelecimento
    r = await client.post("/api/v1/reports/exports", headers=so.headers, json={
        "report_key": "downloads", "format": "csv", "audience": "ESTABELECIMENTO", "org_id": str(w["est_org"].id), **q})
    assert r.status_code == 201, r.text
    content = (await client.get(r.json()["download_url"], headers=so.headers)).content.decode("utf-8-sig")
    parsed = list(csv.reader(io.StringIO(content), delimiter=";"))
    assert parsed[0][:3] == ["Tela", "Nome", "Downloads iniciados"] and len(parsed) == 3
    assert other_screen.friendly_id not in content


async def test_dashboards_common_period_filter(client, db, storage, mail):
    w = await _world(db, client)
    today = local_date(utcnow(), "America/Sao_Paulo")
    old_day = local_date(w["old"], "America/Sao_Paulo")
    old_q = {"start": (old_day - timedelta(days=1)).isoformat(), "end": old_day.isoformat()}
    sadm, so, sa = w["s"]["admin"], w["s"]["owner"], w["s"]["adv_a"]

    # padrão (sem parâmetros): mês corrente até hoje, campos antigos intactos
    d = (await client.get("/api/v1/dashboards/admin", headers=sadm.headers)).json()
    assert d["period"] == {**d["period"], "start": today.replace(day=1).isoformat(), "end": today.isoformat(),
                           "explicit": False, "timezone": "America/Sao_Paulo"}
    assert d["kpis"]["exhibitions_today"] == 18 and d["kpis"]["exhibitions_period"] == d["kpis"]["exhibitions_month"]
    assert d["kpis"]["advertisers_exhibited_period"] == 2

    # período antigo: só as 7 exibições preservadas no agregado (tela 2, Campanha Alfa)
    r = await client.get("/api/v1/dashboards/admin", headers=sadm.headers, params=old_q)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["period"]["explicit"] is True and d["period"]["start"] == old_q["start"]
    assert d["kpis"]["exhibitions_period"] == 7 and d["kpis"]["advertisers_exhibited_period"] == 1
    assert d["kpis"]["exhibitions_today"] == 18  # estado/“hoje” não muda com o período
    assert d["kpis"]["screens_online"] == 1 and d["operations"]["campaigns_awaiting_approval"] == 1
    assert d["distribution_by_state"] == [{"uf": "GO", "screens": 2, "online": 1, "exhibitions_period": 7}]
    assert d["operations"]["media_errors_period"] == 0

    r = await client.get("/api/v1/dashboards/establishment", headers=so.headers,
                         params={**old_q, "org_id": str(w["est_org"].id)})
    e = r.json()
    assert e["kpis"]["exhibitions_period"] == 7 and e["kpis"]["external_exhibitions_period"] == 7
    assert e["kpis"]["internal_exhibitions_period"] == 0 and e["kpis"]["external_exhibitions_month"] == 12
    r = await client.get("/api/v1/dashboards/advertiser", headers=sa.headers,
                         params={**old_q, "org_id": str(w["adv_a"].id)})
    a = r.json()
    assert a["period"]["explicit"] is True and a["regions_exhibited"] == [{"uf": "GO", "screens": 1}]
    # mesmo período para o anunciante B: nada exibido
    r = await client.get("/api/v1/dashboards/advertiser", headers=w["s"]["adv_b"].headers,
                         params={**old_q, "org_id": str(w["adv_b"].id)})
    assert r.json()["regions_exhibited"] == []
    # validação do período
    bad = await client.get("/api/v1/dashboards/admin", headers=sadm.headers,
                           params={"start": today.isoformat(), "end": (today - timedelta(days=1)).isoformat()})
    assert bad.status_code == 400 and bad.json()["error"]["code"] == "PERIOD_INVALID"
    big = await client.get("/api/v1/dashboards/admin", headers=sadm.headers,
                           params={"start": (today - timedelta(days=400)).isoformat(), "end": today.isoformat()})
    assert big.status_code == 400 and big.json()["error"]["code"] == "PERIOD_TOO_LONG"


async def test_admin_finance_projection_borders_by_period(client, db, storage, mail):
    """Projeção linear enquanto início ≤ hoje ≤ fim (inclui o ÚLTIMO dia); encerrado só se fim < hoje."""
    admin = await make_user(db, role="admin", mfa=True)  # finance.view
    s = await login_as(client, admin)
    today = local_date(utcnow(), "America/Sao_Paulo")

    async def method(start, end):
        r = await client.get("/api/v1/dashboards/admin", headers=s.headers,
                             params={"start": start.isoformat(), "end": end.isoformat()})
        assert r.status_code == 200, r.text
        return r.json()["finance"]["projection_method"]

    assert (await method(today - timedelta(days=2), today)).startswith("Linear: receita reconhecida em 3 de 3 dias")
    assert (await method(today, today)).startswith("Linear: receita reconhecida em 1 de 1 dias")
    assert (await method(today - timedelta(days=1), today + timedelta(days=1))).startswith("Linear: receita reconhecida em 2 de 3")
    assert (await method(today - timedelta(days=3), today - timedelta(days=1))).startswith("Período encerrado")
    assert (await method(today + timedelta(days=1), today + timedelta(days=2))).startswith("Período futuro")
